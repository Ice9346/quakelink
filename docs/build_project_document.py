from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "docs" / "DEVELOPMENT_DOCUMENTATION.md"
OUTPUT_DOCX = ROOT / "docs" / "QuakeLink_Project_Documentation.docx"
LOGO = ROOT / "assets" / "logos" / "quakelink.png"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 32, 44)
MUTED = RGBColor(95, 108, 120)
LIGHT_FILL = "E8EEF5"
SOFT_FILL = "F4F6F9"
BORDER = "D9E2EC"
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Inches(width)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 8),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["List Bullet"].font.name = "Calibri"
    styles["List Bullet"].font.size = Pt(10.3)
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.38)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.19)
    styles["List Bullet"].paragraph_format.space_after = Pt(3)


def add_header_footer(section):
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("QuakeLink Project Documentation")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    add_field(footer, "PAGE")


def add_cover(doc: Document):
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(str(LOGO), width=Inches(1.1))
        except Exception:
            pass

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(36)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("QUAKELINK")
    set_run_font(run, size=30, color=INK, bold=True)
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Project Development Document")
    set_run_font(run, size=18, color=BLUE, bold=True)
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Frontend | Backend | Database")
    set_run_font(run, size=12, color=MUTED, bold=True)
    p.paragraph_format.space_after = Pt(36)

    table = doc.add_table(rows=4, cols=2)
    set_table_width(table, [1.7, 4.6])
    rows = [
        ("Document Version", "1.0.0"),
        ("Project Path", "/Users/bae/quakelink_app"),
        ("Prepared For", "QuakeLink application development and maintenance"),
        ("Date", date(2026, 7, 5).isoformat()),
    ]
    for row, (label, value) in zip(table.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell, top=90, bottom=90)
            if idx == 0:
                set_cell_shading(cell, LIGHT_FILL)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    set_run_font(run, size=10, color=INK, bold=(idx == 0))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "เอกสารเล่มนี้สรุปภาพรวมระบบ สถาปัตยกรรม โครงสร้างไฟล์ API "
        "WebSocket และ schema ฐานข้อมูลจากโปรเจกต์ปัจจุบัน"
    )
    set_run_font(run, size=10.5, color=MUTED)
    doc.add_page_break()


def add_overview(doc: Document):
    doc.add_heading("Document Overview", level=1)
    p = doc.add_paragraph(
        "เอกสารนี้จัดรูปแบบจากเอกสารพัฒนา QuakeLink ฉบับเต็ม "
        "เพื่อใช้เป็นเล่มอ้างอิงสำหรับการพัฒนา การส่งต่อโปรเจกต์ และการดูแลระบบในอนาคต"
    )
    p.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=3)
    set_table_width(table, [2.15, 2.15, 2.15])
    headers = ["Frontend", "Backend", "Database"]
    descriptions = [
        "Flutter app, routing, services, realtime store, UI components",
        "Node.js Express API, WebSocket, ThingsBoard bridge, auth, OpenAI",
        "PostgreSQL, Prisma models, relationships, migration, data flow",
    ]
    for idx, cell in enumerate(table.rows[0].cells):
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_borders(cell)
        set_cell_margins(cell, top=120, bottom=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(headers[idx])
        set_run_font(r, size=12, color=DARK_BLUE, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(descriptions[idx])
        set_run_font(r2, size=9.5, color=INK)

    doc.add_heading("สารบัญแบบย่อ", level=1)
    for item in [
        "1. Project Overview",
        "2. Tech Stack Summary",
        "3. Frontend",
        "4. Backend",
        "5. Database",
        "6. Running The Project",
        "7. Verification Commands",
        "8. Development Notes",
        "9. Quick Reference",
        "10. Document Version",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)

    doc.add_page_break()


def add_code_block(doc: Document, lines: list[str]):
    if not lines:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    shade_paragraph(p, SOFT_FILL)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=8.2, color=INK)


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "")
    return text.strip()


def add_markdown_content(doc: Document, markdown: str):
    lines = markdown.splitlines()
    in_code = False
    code_lines: list[str] = []
    skip_first_title = True

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = clean_inline(heading.group(2))
            if level == 1 and skip_first_title:
                skip_first_title = False
                continue
            if level == 2 and re.match(r"\d+\.", text):
                doc.add_page_break()
            mapped = 1 if level <= 2 else 2 if level == 3 else 3
            doc.add_heading(text, level=mapped)
            continue

        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(clean_inline(bullet.group(1)))
            continue

        if re.match(r"^\s*\d+\.\s+.+$", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(clean_inline(re.sub(r"^\s*\d+\.\s+", "", line)))
            continue

        p = doc.add_paragraph()
        p.add_run(clean_inline(line))


def build():
    doc = Document()
    configure_document(doc)
    add_header_footer(doc.sections[0])
    add_cover(doc)
    add_overview(doc)
    add_markdown_content(doc, SOURCE_MD.read_text(encoding="utf-8"))

    for section in doc.sections:
        add_header_footer(section)

    doc.core_properties.title = "QuakeLink Project Development Document"
    doc.core_properties.subject = "Frontend, Backend, and Database documentation"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.keywords = "QuakeLink, Flutter, Node.js, PostgreSQL, Prisma"
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    print(build())
